"""Exportación de cualquier tabla (Kardex o reportes) a Excel, PDF, CSV o TXT."""
import csv
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas as _canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def exportar(fmt, titulo, enc, filas, pie=""):
    buf = io.BytesIO()
    if fmt in ("csv", "txt"):
        s = io.StringIO()
        w = csv.writer(s, delimiter="," if fmt == "csv" else "\t")
        w.writerow(enc)
        w.writerows(filas)
        buf.write(s.getvalue().encode("utf-8-sig"))
    elif fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Datos"
        ws.append([titulo])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(enc)
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", start_color="1769AA")
        for f in filas:
            ws.append(f)
        if pie:
            ws.append([pie])
        wb.save(buf)
    elif fmt == "pdf":
        datos = [enc] + [["" if v is None else (f"{v:,.2f}" if isinstance(v, float) else str(v)) for v in f] for f in filas]
        t = Table(datos, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1769AA")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
        est = getSampleStyleSheet()
        SimpleDocTemplate(buf, pagesize=landscape(letter) if len(enc) > 5 else letter).build(
            [Paragraph(titulo, est["Title"]), t] + ([Spacer(1, 8), Paragraph(pie, est["Normal"])] if pie else []))
    else:
        raise ValueError(fmt)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Kardex con el mismo formato de la pantalla: Entradas (verde), Salidas (rojo) y Saldo (azul),
# cada bloque con Cant., Precio y Total.
# ---------------------------------------------------------------------------
VERDE, ROJO, AZUL, MARINO = "#16A085", "#E74C3C", "#1769AA", "#0F2747"
OSCURO = {VERDE: "#118A73", ROJO: "#C9402F", AZUL: "#125A92"}
TINTE = {VERDE: "#F1FBF8", ROJO: "#FFF5F4", AZUL: "#F2F7FC"}
GRUPOS = (("ENTRADAS", VERDE, 2), ("SALIDAS", ROJO, 5), ("SALDO", AZUL, 8))


def _cant(v):
    return f"{int(v):,}" if float(v) == int(v) else f"{v:,.2f}".rstrip("0").rstrip(".")


def _periodo(ini, fin):
    return "todo el historial" if not ini and not fin else f"{ini or 'inicio'} a {fin or 'hoy'}"


def _bloques(f, mon):
    """Las 9 celdas numéricas de una fila: entradas, salidas y saldo (None si ese bloque no aplica)."""
    din = lambda v: f"{mon}{(round(v, 2) or 0.0):,.2f}"
    def b(x):
        return [_cant(x[0]), din(x[1]), din(x[2])] if x else None
    return b(f["e"]), b(f["s"]), [_cant(f["v"][0]), din(f["v"][1]), din(f["v"][2])]


def _totales(filas):
    ce = sum(f["e"][0] for f in filas if f["e"]); te = sum(f["e"][2] for f in filas if f["e"])
    cs = sum(f["s"][0] for f in filas if f["s"]); ts = sum(f["s"][2] for f in filas if f["s"])
    return ce, te, cs, ts, (filas[-1]["v"] if filas else (0, 0, 0))


def _numerado(pie):
    class Numerado(_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            for estado in self._paginas:
                self.__dict__.update(estado)
                self.setFont("Helvetica", 8)
                self.setFillColor(colors.HexColor("#6B7280"))
                self.drawString(28, 16, pie)
                self.drawRightString(self._pagesize[0] - 28, 16, f"Página {self._pageNumber} de {total}")
                super().showPage()
            super().save()
    return Numerado


def _kardex_pdf(titulo, filas, mon, empresa, metodo, ini, fin, usuario):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=28, rightMargin=28, topMargin=30, bottomMargin=34, title=titulo)
    est = getSampleStyleSheet()
    t_est = ParagraphStyle("t", parent=est["Title"], fontSize=17, spaceAfter=3)
    s_est = ParagraphStyle("s", parent=est["Normal"], fontSize=9, textColor=colors.HexColor("#4B5563"), alignment=TA_CENTER, spaceAfter=12)
    d_est = ParagraphStyle("d", fontName="Helvetica", fontSize=8.5, leading=10)
    c_est = ParagraphStyle("c", fontName="Helvetica-Oblique", fontSize=7.5, leading=9, textColor=colors.HexColor("#6B7280"))
    story = [Paragraph(titulo, t_est), Paragraph(f"{empresa} · Método de valuación: <b>{metodo}</b> · Período: {_periodo(ini, fin)}", s_est)]
    if not filas:
        story.append(Paragraph("No hay movimientos en el período seleccionado.", est["Normal"]))
        doc.build(story, canvasmaker=_numerado(f"{empresa} · emitido el {date.today().isoformat()} por {usuario}"))
        buf.seek(0)
        return buf
    datos = [["Fecha", "Documento", "ENTRADAS", "", "", "SALIDAS", "", "", "SALDO", "", ""], ["", ""] + ["Cant.", "Precio", "Total"] * 3]
    ts = [("SPAN", (0, 0), (0, 1)), ("SPAN", (1, 0), (1, 1)), ("BACKGROUND", (0, 0), (1, 1), colors.HexColor(MARINO))]
    for nombre, col, c0 in GRUPOS:
        ts += [("SPAN", (c0, 0), (c0 + 2, 0)), ("BACKGROUND", (c0, 0), (c0 + 2, 0), colors.HexColor(col)), ("BACKGROUND", (c0, 1), (c0 + 2, 1), colors.HexColor(OSCURO[col])),
               ("BACKGROUND", (c0, 2), (c0 + 2, -1), colors.HexColor(TINTE[col]))]
    vacios = []
    for f in filas:
        r = len(datos)
        ent, sal, sld = _bloques(f, mon)
        datos.append([f["fecha"], Paragraph("(continúa)", c_est) if f["cont"] else Paragraph(f["doc"] or "–", d_est)] + (ent or ["–", "", ""]) + (sal or ["–", "", ""]) + sld)
        for bloque, c0 in ((ent, 2), (sal, 5)):
            if bloque is None:
                ts.append(("SPAN", (c0, r), (c0 + 2, r)))
                vacios.append(("ALIGN", (c0, r), (c0 + 2, r), "CENTER"))
                vacios.append(("TEXTCOLOR", (c0, r), (c0 + 2, r), colors.HexColor("#9CA3AF")))
    ce, te, cs, ts_, v = _totales(filas)
    r = len(datos)
    din = lambda x: f"{mon}{(round(x, 2) or 0.0):,.2f}"
    datos.append(["TOTALES", "", _cant(ce), "", din(te), _cant(cs), "", din(ts_), _cant(v[0]), din(v[1]), din(v[2])])
    ts += [("SPAN", (0, r), (1, r)), ("FONTNAME", (0, r), (-1, r), "Helvetica-Bold"), ("LINEABOVE", (0, r), (-1, r), 1.2, colors.HexColor(MARINO)), ("BACKGROUND", (0, r), (1, r), colors.HexColor("#E8EDF3"))]
    ts += [("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"), ("TEXTCOLOR", (0, 0), (-1, 1), colors.white), ("ALIGN", (0, 0), (-1, 1), "CENTER"), ("FONTSIZE", (0, 0), (-1, 1), 9),
           ("FONTSIZE", (0, 2), (-1, -1), 8.5), ("ALIGN", (2, 2), (-1, -1), "RIGHT"), ("ALIGN", (0, 2), (1, -1), "LEFT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5DCE6")), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
           ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5)]
    tabla = Table(datos, colWidths=[58, 78] + [48, 64, 76] * 3, repeatRows=2)
    tabla.setStyle(TableStyle(ts + vacios))
    story += [tabla, Spacer(1, 10), Paragraph(f"<b>Existencia final:</b> {_cant(v[0])} &nbsp;&nbsp; <b>Precio:</b> {din(v[1])} &nbsp;&nbsp; <b>Valor del inventario:</b> {din(v[2])}", est["Normal"])]
    doc.build(story, canvasmaker=_numerado(f"{empresa} · emitido el {date.today().isoformat()} por {usuario}"))
    buf.seek(0)
    return buf


def _kardex_xlsx(titulo, filas, mon, empresa, metodo, ini, fin, usuario):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Kardex"
    ws.sheet_view.showGridLines = False
    lado = Side(style="thin", color="D5DCE6")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
    ws.merge_cells("A1:K1")
    ws["A1"] = titulo
    ws["A1"].font = Font(name="Calibri", bold=True, size=15, color=MARINO[1:])
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:K2")
    ws["A2"] = f"{empresa} · Método de valuación: {metodo} · Período: {_periodo(ini, fin)} · Emitido el {date.today().isoformat()} por {usuario}"
    ws["A2"].font = Font(name="Calibri", size=10, color="4B5563")
    ws["A2"].alignment = Alignment(horizontal="center")
    blanco = Font(name="Calibri", bold=True, color="FFFFFF")
    centro = Alignment(horizontal="center", vertical="center")
    for col, texto in ((1, "Fecha"), (2, "Documento")):
        ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
        c = ws.cell(4, col, texto)
        c.font, c.alignment = blanco, centro
        for r in (4, 5):
            ws.cell(r, col).fill = PatternFill("solid", start_color=MARINO[1:])
            ws.cell(r, col).border = borde
    for nombre, col, c0 in GRUPOS:
        ws.merge_cells(start_row=4, start_column=c0 + 1, end_row=4, end_column=c0 + 3)
        ws.cell(4, c0 + 1, nombre).font, ws.cell(4, c0 + 1).alignment = Font(name="Calibri", bold=True, color="FFFFFF", size=11), centro
        for k, t in enumerate(("Cant.", "Precio", "Total")):
            ws.cell(4, c0 + 1 + k).fill = PatternFill("solid", start_color=col[1:])
            ws.cell(4, c0 + 1 + k).border = borde
            s = ws.cell(5, c0 + 1 + k, t)
            s.font, s.alignment, s.border = blanco, centro, borde
            s.fill = PatternFill("solid", start_color=OSCURO[col][1:])
    fmt_c, fmt_d = "#,##0.00", f'"{mon}"#,##0.00'
    r = 6
    for f in filas:
        ws.cell(r, 1, f["fecha"])
        ws.cell(r, 2, "(continúa)" if f["cont"] else (f["doc"] or "–"))
        if f["cont"]:
            ws.cell(r, 2).font = Font(name="Calibri", italic=True, color="6B7280", size=9)
        for bloque, c0, col in ((f["e"], 3, VERDE), (f["s"], 6, ROJO), (f["v"], 9, AZUL)):
            for k in range(3):
                c = ws.cell(r, c0 + k, (bloque[k] + 0.0) if bloque else None)
                c.number_format = fmt_c if k == 0 else fmt_d
                c.fill = PatternFill("solid", start_color=TINTE[col][1:])
                c.alignment = Alignment(horizontal="right")
        for k in range(1, 12):
            ws.cell(r, k).border = borde
        r += 1
    if filas:
        ce, te, cs, ts_, v = _totales(filas)
        for k, val in {1: "TOTALES", 3: ce, 5: te, 6: cs, 8: ts_, 9: v[0], 10: v[1], 11: v[2]}.items():
            c = ws.cell(r, k, val + 0.0 if k > 1 else val)
            c.font = Font(name="Calibri", bold=True)
            c.fill = PatternFill("solid", start_color="E8EDF3")
            c.border = Border(left=lado, right=lado, bottom=lado, top=Side(style="medium", color=MARINO[1:]))
            if k in (3, 6, 9):
                c.number_format = fmt_c
            elif k > 3:
                c.number_format = fmt_d
        ws.cell(r, 2).fill = PatternFill("solid", start_color="E8EDF3")
        ws.cell(r, 2).border = Border(top=Side(style="medium", color=MARINO[1:]))
        ws.cell(r + 2, 1, f"Existencia final: {_cant(v[0])}   ·   Precio: {mon}{v[1]:,.2f}   ·   Valor del inventario: {mon}{v[2]:,.2f}").font = Font(name="Calibri", bold=True, color=MARINO[1:])
    else:
        ws.cell(6, 1, "No hay movimientos en el período seleccionado.")
    for i, w in enumerate((12, 14, 11, 13, 15, 11, 13, 15, 11, 13, 15), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C6"
    ws.print_title_rows = "4:5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def exportar_kardex(fmt, titulo, filas, mon="Q", empresa="", metodo="", ini="", fin="", usuario=""):
    """Kardex en PDF o Excel con el formato de la pantalla (bloques de colores, subtítulos y totales)."""
    if fmt == "pdf":
        return _kardex_pdf(titulo, filas, mon, empresa, metodo, ini, fin, usuario)
    if fmt == "xlsx":
        return _kardex_xlsx(titulo, filas, mon, empresa, metodo, ini, fin, usuario)
    raise ValueError(fmt)
