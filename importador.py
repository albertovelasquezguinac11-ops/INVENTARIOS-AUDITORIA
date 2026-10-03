"""Importador de archivos para el sistema de inventarios.

Detecta el formato del archivo subido (imagen, PDF, Word, Excel/CSV o texto)
y extrae de él una lista de movimientos: fecha, tipo, cantidad, costo_unitario.

Para Excel/CSV los datos ya vienen estructurados por columnas, así que la
lectura es directa y confiable. Para fotografías, PDF escaneado y Word se usa
texto libre (OCR o extracción de texto) y luego se interpreta línea por línea
con expresiones regulares — por eso el resultado SIEMPRE se muestra al
usuario en una pantalla de revisión antes de guardarse: el OCR nunca es
100% exacto y el sistema no debe fallar por confiar ciegamente en él.
"""
import csv
import datetime as _dt
import re
import unicodedata
from typing import Any, Dict, List, Optional, Tuple, cast

import pandas as pd
import pdfplumber
import pytesseract  # pyright: ignore[reportMissingTypeStubs]
from PIL import Image
from docx import Document

# pytesseract y pdf2image no publican type stubs (.pyi), por eso pyright/Pylance
# no puede inferir sus tipos de retorno con precisión. Se envuelven en funciones
# pequeñas que fuerzan el tipo real que sabemos que devuelven en modo texto.


def _ocr_imagen(imagen: Any) -> str:
    # pytesseract no publica stubs de tipos (.pyi), así que Pylance no puede saber
    # que en modo texto siempre devuelve str; se lo indicamos explícitamente aquí.
    import os, shutil
    if not shutil.which("tesseract"):  # en Windows suele instalarse aquí y no queda en el PATH
        for ruta in (r"C:\Program Files\Tesseract-OCR\tesseract.exe", r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
            if os.path.exists(ruta):
                pytesseract.pytesseract.tesseract_cmd = ruta  # pyright: ignore
                break
    try:
        return cast(str, pytesseract.image_to_string(imagen, lang="spa"))  # pyright: ignore[reportUnknownMemberType]
    except pytesseract.TesseractNotFoundError:  # pyright: ignore
        raise ValueError("para leer fotos y PDF escaneados hay que instalar Tesseract-OCR en la computadora")
    except pytesseract.TesseractError:  # pyright: ignore  (falta el idioma español: se usa el inglés)
        return cast(str, pytesseract.image_to_string(imagen, lang="eng"))  # pyright: ignore

FilaImportada = Dict[str, Any]

FORMATOS_IMAGEN = ("png", "jpg", "jpeg", "bmp", "webp", "tiff")

PATRON_FECHA = re.compile(r"\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{4}")
PATRON_TIPO = re.compile(r"\b(entrada|salida|ent|sal)\b", re.IGNORECASE)
PATRON_NUMERO = re.compile(r"\d+(?:[.,]\d+)?")


def _normalizar_fecha(texto: str) -> str:
    texto = texto.strip().replace("/", "-")
    partes = texto.split("-")
    if len(partes[0]) == 4:
        return texto
    dia, mes, anio = partes
    return f"{anio}-{mes.zfill(2)}-{dia.zfill(2)}"


def _normalizar_tipo(texto: str) -> str:
    return "entrada" if texto.strip().lower().startswith("ent") else "salida"


def _a_float_opcional(valor: Any) -> Optional[float]:
    if valor is None:
        return None
    texto = str(valor).strip().lower()
    if texto in ("", "nan", "none"):
        return None
    return float(valor)


def _parsear_linea(linea: str) -> Optional[FilaImportada]:
    """Extrae fecha, tipo, cantidad y costo unitario de una línea de texto libre.
    Devuelve None si la línea no trae lo mínimo indispensable (fecha + tipo + cantidad)."""
    fecha_m = PATRON_FECHA.search(linea)
    tipo_m = PATRON_TIPO.search(linea)
    if not fecha_m or not tipo_m:
        return None

    resto = linea[fecha_m.end():]
    numeros = PATRON_NUMERO.findall(resto)
    if not numeros:
        return None

    cantidad = float(numeros[0].replace(",", ""))
    costo_unitario = float(numeros[1].replace(",", "")) if len(numeros) > 1 else None

    return {
        "codigo_producto": None,
        "nombre_producto": None,
        "fecha": _normalizar_fecha(fecha_m.group(0)),
        "tipo": _normalizar_tipo(tipo_m.group(0)),
        "cantidad": cantidad,
        "costo_unitario": costo_unitario,
    }


def procesar_texto_libre(texto: str) -> List[FilaImportada]:
    filas: List[FilaImportada] = []
    for linea in texto.splitlines():
        linea = linea.strip()
        if not linea:
            continue
        fila = _parsear_linea(linea)
        if fila:
            filas.append(fila)
    return filas


def procesar_imagen(ruta: str) -> Tuple[List[FilaImportada], str]:
    imagen = Image.open(ruta)
    texto = _ocr_imagen(imagen)
    return procesar_texto_libre(texto), texto


def procesar_pdf(ruta: str) -> Tuple[List[FilaImportada], str]:
    texto_total = ""
    with pdfplumber.open(ruta) as pdf:
        for pagina in pdf.pages:
            texto_total += (pagina.extract_text() or "") + "\n"

    if texto_total.strip():
        return procesar_texto_libre(texto_total), texto_total

    # El PDF no trae texto seleccionable (es un escaneo) -> se rasteriza y se aplica OCR
    from pdf2image import convert_from_path  # pyright: ignore[reportUnknownVariableType]
    paginas_imagen = cast(List[Any], convert_from_path(ruta))
    texto_total = ""
    for imagen_pagina in paginas_imagen:
        texto_total += _ocr_imagen(imagen_pagina) + "\n"
    return procesar_texto_libre(texto_total), texto_total


def procesar_docx(ruta: str) -> Tuple[List[FilaImportada], str]:
    doc = Document(ruta)
    filas: List[FilaImportada] = []
    texto_total = ""

    for tabla in doc.tables:
        for fila_tabla in tabla.rows[1:]:  # se asume que la primera fila es encabezado
            linea = " ".join(c.text.strip() for c in fila_tabla.cells)
            texto_total += linea + "\n"
            fila = _parsear_linea(linea)
            if fila:
                filas.append(fila)

    if not filas:
        for parrafo in doc.paragraphs:
            texto_total += parrafo.text + "\n"
        filas = procesar_texto_libre(texto_total)

    return filas, texto_total


def procesar_excel_o_csv(ruta: str, extension: str) -> Tuple[List[FilaImportada], None]:
    df: pd.DataFrame = pd.read_csv(ruta) if extension == "csv" else pd.read_excel(ruta)  # pyright: ignore[reportUnknownMemberType]
    df.columns = [str(c).strip().lower() for c in df.columns]

    filas: List[FilaImportada] = []
    for _, fila_serie in df.iterrows():
        fila: Dict[str, Any] = {str(k): v for k, v in fila_serie.to_dict().items()}
        filas.append({
            "codigo_producto": str(fila.get("codigo", "")).strip() or None,
            "nombre_producto": str(fila.get("nombre", "")).strip() or None,
            "fecha": _normalizar_fecha(str(fila.get("fecha", ""))),
            "tipo": _normalizar_tipo(str(fila.get("tipo", "entrada"))),
            "cantidad": _a_float_opcional(fila.get("cantidad")) or 0.0,
            "costo_unitario": _a_float_opcional(fila.get("costo_unitario")),
        })
    return filas, None


def procesar_archivo(ruta: str, extension: str) -> Tuple[List[FilaImportada], Optional[str]]:
    """Punto de entrada único. Devuelve (filas, texto_crudo).
    texto_crudo es None cuando el origen ya era una tabla estructurada (Excel/CSV)."""
    extension = extension.lower().lstrip(".")

    if extension in ("xlsx", "xls", "csv"):
        return procesar_excel_o_csv(ruta, extension)
    if extension == "docx":
        return procesar_docx(ruta)
    if extension == "pdf":
        return procesar_pdf(ruta)
    if extension in FORMATOS_IMAGEN:
        return procesar_imagen(ruta)
    if extension == "txt":
        with open(ruta, encoding="utf-8", errors="ignore") as f:
            texto = f.read()
        return procesar_texto_libre(texto), texto

    raise ValueError(f"El formato .{extension} todavía no está soportado.")


# ---------------------------------------------------------------------------
# Importación flexible: detección de columnas por sinónimos y por contenido
# ---------------------------------------------------------------------------
SIN = {
    "codigo": ["codigo", "cod", "sku", "clave", "id", "referencia", "ref"],
    "nombre": ["nombre", "producto", "articulo", "descripcion", "item", "detalle", "mercaderia"],
    "fecha": ["fecha", "dia", "date", "f"],
    "tipo": ["tipo", "movimiento", "mov", "operacion", "concepto", "transaccion"],
    "cantidad": ["cantidad", "cant", "unidades", "uds", "qty", "piezas"],
    "costo": ["costo", "precio", "p unit", "p unitario", "valor unitario", "c u", "unitario"],
}


def _n(t: Any) -> str:
    t = unicodedata.normalize("NFD", str(t)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", t)).strip()


_SIN = {k: [_n(x) for x in v] for k, v in SIN.items()}


def _coincide(h: str, sins: List[str]) -> bool:
    return bool(h) and any(h == s or s in h.split() or (len(s) > 4 and s in h) for s in sins)


def _es_enc(c: Any) -> bool:
    h = _n(c)
    return any(_coincide(h, v) for v in _SIN.values())


def _tipo(v: Any) -> Optional[str]:
    t = _n(v)
    if t in ("e", "in", "ingreso", "compra") or t.startswith("entr"):
        return "entrada"
    if t in ("s", "out", "egreso", "venta") or t.startswith("sali"):
        return "salida"
    return None


def _fecha(v: Any) -> Optional[str]:
    s = str(v).strip()
    if re.match(r"\d{4}-\d{2}-\d{2}", s):
        s = s[:10]
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y", "%d-%m-%y"):
        try:
            return _dt.datetime.strptime(s, f).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def _num2(v: Any) -> Optional[float]:
    s = re.sub(r"[Qq$\s]", "", str(v))
    if s in ("", "nan", "None"):
        return None
    s = s.replace(",", "") if re.search(r"\d,\d{3}", s) else s.replace(",", ".")
    return float(s)


def _dividir(linea: str) -> List[str]:
    p = [x.strip() for x in re.split(r"\t|;|\||\s{2,}", linea) if x.strip()]
    return p if len(p) >= 3 else linea.split()


def _tabla(m: List[List[Any]]) -> Tuple[List[str], List[List[str]]]:
    m = [[("" if c is None else str(c).strip()) for c in r] for r in m if any(str(c or "").strip() for c in r)]
    if not m:
        return [], []
    hi = next((k for k, r in enumerate(m[:6]) if sum(_es_enc(c) for c in r) >= 2), None)
    n = max(len(r) for r in m)
    enc = m[hi] if hi is not None else []
    enc = enc + [f"Columna {k + 1}" for k in range(len(enc), n)]
    filas = [r for r in m[(hi + 1 if hi is not None else 0):] if r != m[hi if hi is not None else 0] or hi is None]
    return enc, filas


def extraer_tabla(ruta: str, ext: str) -> Tuple[List[str], List[List[str]], Optional[str]]:
    """Convierte CUALQUIER archivo soportado en una tabla cruda (encabezados + filas de texto)."""
    ext = ext.lower().lstrip(".")
    texto, matriz = "", None
    if ext in ("xlsx", "xls"):
        matriz = pd.read_excel(ruta, header=None, dtype=str).fillna("").values.tolist()  # pyright: ignore
    elif ext == "csv":
        with open(ruta, encoding="utf-8-sig", errors="ignore") as fh:
            muestra = fh.read(4096)
            fh.seek(0)
            try:
                dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
            except csv.Error:
                dialecto = csv.excel
            matriz = list(csv.reader(fh, dialecto))
    elif ext == "docx":
        doc = Document(ruta)
        tablas = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
        if tablas:
            matriz = max(tablas, key=len)
        else:
            texto = "\n".join(p.text for p in doc.paragraphs)
    elif ext == "pdf":
        with pdfplumber.open(ruta) as pdf:
            tablas = [t for pg in pdf.pages for t in pg.extract_tables()]
            texto = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
        if tablas:
            matriz = [r for t in tablas for r in t]
        elif not texto.strip():
            from pdf2image import convert_from_path  # pyright: ignore
            texto = "\n".join(_ocr_imagen(im) for im in cast(List[Any], convert_from_path(ruta)))
    elif ext in FORMATOS_IMAGEN:
        texto = _ocr_imagen(Image.open(ruta))
    elif ext == "txt":
        with open(ruta, encoding="utf-8", errors="ignore") as fh:
            texto = fh.read()
    else:
        raise ValueError(f"El formato .{ext} todavía no está soportado.")
    if matriz is None:
        enc, filas = tabla_desde_texto(texto)
    else:
        enc, filas = _tabla(matriz)
    return enc, filas, (texto or None)


def detectar(enc: List[str], filas: List[List[str]]) -> Dict[str, int]:
    """Adivina qué columna es cada campo: primero por el nombre del encabezado (sinónimos), luego por el contenido."""
    m: Dict[str, int] = {}
    usados: set = set()
    hs = [_n(h) for h in enc]
    for campo, sins in _SIN.items():
        for k, h in enumerate(hs):
            if k not in usados and _coincide(h, sins):
                m[campo] = k
                usados.add(k)
                break

    def vals(k: int) -> List[str]:
        return [str(r[k]).strip() for r in filas[:25] if k < len(r) and str(r[k]).strip()]

    libres = [k for k in range(len(enc)) if k not in usados]
    for campo, prueba in (("fecha", lambda v: _fecha(v) is not None), ("tipo", lambda v: _tipo(v) is not None)):
        if campo not in m:
            for k in libres:
                vs = vals(k)
                if vs and sum(prueba(v) for v in vs) / len(vs) >= 0.6:
                    m[campo] = k
                    libres.remove(k)
                    break
    nums = [k for k in libres if vals(k) and sum(bool(re.fullmatch(r"[Qq$]?\s*-?[\d.,]+", v)) for v in vals(k)) / len(vals(k)) >= 0.6]  # tolera alguna celda mala
    if "cantidad" not in m and nums:
        m["cantidad"] = nums.pop(0)
    if "costo" not in m and nums:
        m["costo"] = nums.pop(0)
    if "nombre" not in m:
        txt = [k for k in libres if k not in m.values() and vals(k) and any(re.search(r"[A-Za-z]", v) for v in vals(k))]
        if txt:
            m["nombre"] = txt[0]
    return m


def construir(filas: List[List[str]], mapa: Dict[str, int], tipo_def: str = "", prod_def: Optional[Tuple[str, str]] = None) -> List[Dict[str, Any]]:
    """Aplica el mapeo de columnas y valida cada fila. Nunca lanza error: marca el problema en 'error'."""
    out: List[Dict[str, Any]] = []
    for r in filas:
        def g(k: str) -> str:
            return str(r[mapa[k]]).strip() if k in mapa and mapa[k] < len(r) else ""
        err: List[str] = []
        f = _fecha(g("fecha"))
        if not f:
            err.append("fecha inválida")
        t = _tipo(g("tipo")) if g("tipo") else (tipo_def or None)
        if not t:
            err.append("tipo no reconocido")
        try:
            c = _num2(g("cantidad"))
        except ValueError:
            c = None
        if not c or c <= 0:
            err.append("cantidad inválida")
        try:
            k2 = _num2(g("costo"))
        except ValueError:
            k2 = None
            err.append("costo inválido")
        if t == "entrada" and not k2:
            err.append("falta el costo")
        cod, nom = g("codigo"), g("nombre")
        if not cod and not nom and prod_def:
            cod, nom = prod_def
        if not cod and not nom:
            err.append("falta el producto")
        out.append({"codigo_producto": cod or None, "nombre_producto": nom or None, "fecha": f or g("fecha"),
                    "tipo": t or "", "cantidad": c, "costo_unitario": k2, "error": "; ".join(err)})
    return out


# ---------------------------------------------------------------------------
# Texto libre: entiende frases como "el 5 de septiembre se vendieron 20 unidades de cuaderno"
# ---------------------------------------------------------------------------
MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9,
         "octubre": 10, "noviembre": 11, "diciembre": 12, "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8,
         "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dic": 12}
_SAL_RE = re.compile(r"\b(vend\w*|venta\w*|sali\w*|despach\w*|sac\w*|entreg\w*|egres\w*)", re.I)
_ENT_RE = re.compile(r"\b(compr\w*|entrad\w*|entr[oó]\w*|ingres\w*|lleg\w*|recib\w*|adquir\w*|abastec\w*|repuest\w*)", re.I)
_PALABRA = r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ][A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9\-]*"
_CORTE = re.compile(r"\s+(?:el|la|los|las|a|con|por|en|para|del|al|y|que|se|c/u|cada|hoy|ayer|vendid\w*|comprad\w*|salier\w*|entrar\w*|recibid\w*)\b", re.I)
_RUIDO = {"se", "el", "la", "los", "las", "de", "del", "un", "una", "dia", "día", "fue", "fueron", "hoy", "ayer", "en", "al", "y", "que", "unidades", "unidad"}


def _fecha_natural(s: str) -> Tuple[Optional[str], str]:
    hoy, t = _dt.date.today(), s
    m = re.search(r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b", t)
    if m:
        y, mo, d = map(int, m.groups())
    else:
        m = re.search(r"\b(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})\b", t)
        if m:
            d, mo, y = map(int, m.groups())
            y = y + 2000 if y < 100 else y
        else:
            m = re.search(r"\b(\d{1,2})\s+(?:de\s+)?([A-Za-záéíóú]{3,10})\.?(?:\s+(?:de|del)\s+(\d{4}))?", t)
            if m and m.group(2).lower() in MESES:
                d, mo, y = int(m.group(1)), MESES[m.group(2).lower()], int(m.group(3)) if m.group(3) else hoy.year
            else:
                m = re.search(r"\b(\d{1,2})/(\d{1,2})\b", t)
                if m:
                    d, mo = map(int, m.groups())
                    y = hoy.year
                else:
                    r = re.search(r"\b(hoy|ayer|anteayer|antier)\b", t, re.I)
                    if not r:
                        return None, s
                    f = hoy - _dt.timedelta(days={"hoy": 0, "ayer": 1, "anteayer": 2, "antier": 2}[r.group(1).lower()])
                    return f.isoformat(), s[:r.start()] + " " + s[r.end():]
    try:
        f = _dt.date(y, mo, d).isoformat()
    except ValueError:
        return None, s
    return f, s[:m.start()] + " " + s[m.end():]


def _precio_natural(s: str) -> Tuple[Optional[float], str]:
    m = re.search(r"(?:\b(?:a|por|precio|costo|valor)\s+(?:de\s+)?)?(?<![A-Za-z])[Qq$]\s*(\d+(?:[.,]\d+)?)", s)
    if not m:
        m = re.search(r"\ba\s+(\d+(?:[.,]\d+)?)\s*(?:c/u|cada\s+(?:uno|una)|la\s+unidad|por\s+unidad)", s, re.I)
    if not m:
        return None, s
    try:
        return _num2(m.group(1)), s[:m.start()] + " " + s[m.end():]
    except ValueError:
        return None, s


def _cantidad_producto(s: str) -> Optional[Tuple[float, str]]:
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*(?:unidades?|uds?\.?|piezas?|pzas?\.?|pz|u\b)?\s*(?:de\s+)?(?:(?:la|el|los|las|un|una)\s+)?(" + _PALABRA + r"(?:\s+" + _PALABRA + r"){0,3})", s)
    if m:
        cant, prod = _num2(m.group(1)), _CORTE.split(m.group(2), 1)[0].strip()
    else:
        m = re.search(r"(\d+(?:[.,]\d+)?)", s)
        if not m:
            return None
        cant = _num2(m.group(1))
        pal = [p for p in re.findall(_PALABRA, s[:m.start()]) if p.lower() not in _RUIDO and not _SAL_RE.match(p) and not _ENT_RE.match(p)]
        prod = " ".join(pal[-3:])
    if not cant or not prod or prod.lower() in _RUIDO:
        return (cant, "") if cant else None
    return cant, prod


def _natural(s: str) -> Optional[List[str]]:
    s = s.strip(" .;")
    if len(s) < 6:
        return None
    sal, ent = _SAL_RE.search(s), _ENT_RE.search(s)
    if not sal and not ent:
        return None
    tipo = "salida" if sal and (not ent or sal.start() < ent.start()) else "entrada"
    f, r = _fecha_natural(s)
    precio, r = _precio_natural(r)
    cp = _cantidad_producto(r)
    if not cp:
        return None
    return [f or "", tipo, cp[1].capitalize(), f"{cp[0]:g}", "" if precio is None else f"{precio:g}"]


def tabla_desde_texto(texto: str) -> Tuple[List[str], List[List[str]]]:
    """Texto plano -> tabla. Si parece una tabla (columnas) la usa tal cual; si son frases, las interpreta una por una."""
    lineas = [l for l in texto.splitlines() if l.strip()]
    enc, filas = _tabla([_dividir(l) for l in lineas])
    m = detectar(enc, filas)
    if "fecha" in m and "cantidad" in m and filas and max(len(r) for r in filas) >= 3:
        return enc, filas
    frases = [x for l in lineas for x in re.split(r"(?<=[a-záéíóú0-9])\.\s+(?=[A-ZÁÉÍÓÚ])|;", l)]
    nat = [r for r in (_natural(x) for x in frases) if r]
    return (["Fecha", "Tipo", "Producto", "Cantidad", "Precio"], nat) if nat else (enc, filas)
